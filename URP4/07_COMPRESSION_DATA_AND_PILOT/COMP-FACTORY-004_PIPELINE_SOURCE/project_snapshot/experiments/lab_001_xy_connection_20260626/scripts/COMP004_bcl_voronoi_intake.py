from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import platform
import re
import sys
import unicodedata
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[3]
LAB = ROOT / "experiments" / "lab_001_xy_connection_20260626"
DEFAULT_CONTRACT = (
    LAB
    / "factories"
    / "COMP-FACTORY-001"
    / "config"
    / "compression_groups_v0_2.json"
)
KST = timezone(timedelta(hours=9))

TARGET_LAYOUT = [
    ("Modulus", 4, 3),
    ("Com. Strength", 6, 0),
    ("Com. strain", 6, 1),
    ("APS", 6, 2),
    ("AS", 6, 3),
    ("Yield strength", 8, 0),
    ("Yield strain", 8, 1),
    ("LCC", 8, 2),
    ("EAS", 8, 3),
    ("Densif. strength", 10, 0),
    ("Densifi. strain", 10, 1),
    ("EAE", 10, 2),
    ("SE", 10, 3),
    ("Total energy", 12, 0),
    ("SEA", 12, 1),
    ("Max. Plateau stress", 12, 2),
]

TARGET_IDS = {
    "Modulus": "COMP-Y-MODULUS",
    "Com. Strength": "COMP-Y-COMP-STRENGTH",
    "Com. strain": "COMP-Y-COMP-STRAIN",
    "APS": "COMP-Y-APS",
    "AS": "COMP-Y-AS",
    "Yield strength": "COMP-Y-YIELD-STRENGTH",
    "Yield strain": "COMP-Y-YIELD-STRAIN",
    "LCC": "COMP-Y-LCC",
    "EAS": "COMP-Y-EAS",
    "Densif. strength": "COMP-Y-DENSIF-STRENGTH",
    "Densifi. strain": "COMP-Y-DENSIFI-STRAIN",
    "EAE": "COMP-Y-EAE",
    "SE": "COMP-Y-SE",
    "Total energy": "COMP-Y-TOTAL-ENERGY",
    "SEA": "COMP-Y-SEA",
    "Max. Plateau stress": "COMP-Y-MAX-PLATEAU-STRESS",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def normalized_text(value: Any) -> str:
    return unicodedata.normalize("NFKC", str(value)).strip()


def as_number(value: Any) -> float | None:
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    return None


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str] | None = None) -> None:
    if fieldnames is None:
        fieldnames = list(rows[0]) if rows else []
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def family_from_key(model_key: str) -> str:
    tail = model_key.split("_", 1)[1] if "_" in model_key else model_key
    if tail.startswith("VF"):
        return "VORONOI"
    match = re.match(r"([A-Z]+)", tail)
    return match.group(1) if match else "UNRESOLVED"


def bcl_key(vf: int, label: Any) -> str:
    return f"VF{vf}_{normalized_text(label)}"


def structural_key(model_id: str) -> str:
    return normalized_text(model_id).split(".", 1)[0].strip()


def voronoi_key(raw_id: Any) -> str | None:
    text = normalized_text(raw_id)
    match = re.fullmatch(r"CP_VF(30|45|60)_C(50|100|200|400|800)(R|U)\.csv", text)
    if not match:
        return None
    vf, radius, orientation = match.groups()
    return f"VF{vf}_VF{vf}_C{radius}_{orientation}"


def extract_bcl_performance(path: Path, vf: int) -> tuple[list[dict], list[dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows: list[dict] = []
    cells: list[dict] = []
    max_blocks = (sheet.max_column - 3) // 12 + 1
    for block_index in range(max_blocks):
        start_col = 3 + 12 * block_index
        label = sheet.cell(2, start_col + 2).value
        if label in (None, ""):
            continue
        label_text = normalized_text(label)
        key = bcl_key(vf, label_text)
        row = {
            "model_key": key,
            "raw_model_id": label_text,
            "source_group": "bcl",
            "family_group": family_from_key(key),
            "vf_group": vf,
            "weight": as_number(sheet.cell(4, start_col).value),
            "measured_vf": as_number(sheet.cell(4, start_col + 1).value),
            "height": as_number(sheet.cell(4, start_col + 2).value),
            "performance_source_file": path.name,
            "performance_sheet": sheet.title,
            "source_block_index": block_index + 1,
        }
        for target, row_number, offset in TARGET_LAYOUT:
            cell = sheet.cell(row_number, start_col + offset)
            value = as_number(cell.value)
            row[f"y__{target}"] = value
            cells.append(
                {
                    "model_key": key,
                    "target_id": TARGET_IDS[target],
                    "target_name": target,
                    "source_file": path.name,
                    "sheet": sheet.title,
                    "cell": f"{get_column_letter(start_col + offset)}{row_number}",
                    "value": value,
                    "finite": value is not None,
                }
            )
        rows.append(row)
    return rows, cells


def extract_voronoi_performance(path: Path, sheet_name: str) -> tuple[list[dict], list[dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    values = list(sheet.iter_rows(values_only=True))
    header = list(values[0])
    target_index = {name: header.index(name) for name, _, _ in TARGET_LAYOUT}
    rows: list[dict] = []
    cells: list[dict] = []
    for excel_row, source_row in enumerate(values[1:], start=2):
        key = voronoi_key(source_row[0])
        if key is None:
            continue
        vf = int(re.match(r"VF(30|45|60)_", key).group(1))
        row = {
            "model_key": key,
            "raw_model_id": source_row[0],
            "source_group": "voronoi",
            "family_group": "VORONOI",
            "vf_group": vf,
            "weight": as_number(source_row[header.index("Weight")]),
            "measured_vf": as_number(source_row[header.index("VF")]),
            "height": as_number(source_row[header.index("Height")]),
            "performance_source_file": path.name,
            "performance_sheet": sheet_name,
            "source_block_index": excel_row - 1,
        }
        for target, _, _ in TARGET_LAYOUT:
            col_index = target_index[target]
            value = as_number(source_row[col_index])
            row[f"y__{target}"] = value
            cells.append(
                {
                    "model_key": key,
                    "target_id": TARGET_IDS[target],
                    "target_name": target,
                    "source_file": path.name,
                    "sheet": sheet_name,
                    "cell": f"{get_column_letter(col_index + 1)}{excel_row}",
                    "value": value,
                    "finite": value is not None,
                }
            )
        rows.append(row)
    return rows, cells


def extract_structural(path: Path, sheet_name: str) -> tuple[list[dict], list[dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    headers = [normalized_text(value) if value is not None else "" for value in next(iterator)]
    feature_headers = headers[1:]
    feature_registry: list[dict] = []
    for index, header in enumerate(feature_headers, start=1):
        role = "structural_descriptor_unselected"
        if header in {"Type", "Source"}:
            role = "excluded_provenance"
        elif header == "Target VF":
            role = "condition_sensitivity"
        elif header == "Actual VF":
            role = "physical_condition_candidate"
        feature_registry.append(
            {
                "feature_id": f"SFX{index:03d}",
                "feature_name_raw": header,
                "column_name": f"x__{header}",
                "feature_role": role,
                "status": "available_unselected" if role != "excluded_provenance" else "excluded",
                "unit": "source_header_or_unresolved",
            }
        )
    rows: list[dict] = []
    for excel_row, source_row in enumerate(iterator, start=2):
        if source_row[0] in (None, ""):
            continue
        model_id = normalized_text(source_row[0])
        key = structural_key(model_id)
        row = {
            "model_key": key,
            "structural_model_id": model_id,
            "source_group": "voronoi" if str(source_row[2]) == "Voronoi_30 mm_merge.xlsx" else "bcl" if str(source_row[2]) == "B, C, L model_30 mm_merge.xlsx" else "other",
            "family_group": family_from_key(key),
            "structural_excel_row": excel_row,
        }
        for header, value in zip(feature_headers, source_row[1:]):
            row[f"x__{header}"] = value
        rows.append(row)
    return rows, feature_registry


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--contract", default=str(DEFAULT_CONTRACT))
    args = parser.parse_args()

    contract_path = Path(args.contract).resolve()
    contract = json.loads(contract_path.read_text(encoding="utf-8"))
    output = ROOT / contract["output_root"] / args.run_id
    if output.exists():
        raise RuntimeError(f"refusing to overwrite existing run: {output}")
    output.mkdir(parents=True)
    started = datetime.now(KST)

    source_audit: list[dict] = []

    def checked(item: dict, role: str) -> Path:
        path = ROOT / item["path"]
        actual = sha256(path)
        expected = item["sha256"].upper()
        source_audit.append(
            {
                "role": role,
                "path": item["path"],
                "bytes": path.stat().st_size,
                "expected_sha256": expected,
                "actual_sha256": actual,
                "status": "PASS" if actual == expected else "FAIL",
            }
        )
        if actual != expected:
            raise RuntimeError(f"source hash mismatch: {path}")
        return path

    structural_item = contract["structural_factors"]
    structural_path = checked(structural_item, "structural_factors")
    structural_rows, feature_registry = extract_structural(
        structural_path, structural_item["sheet"]
    )

    performance_rows: list[dict] = []
    cell_lineage: list[dict] = []
    for item in contract["bcl"]["primary_performance_files"]:
        path = checked(item, f"bcl_primary_vf{item['vf']}")
        rows, cells = extract_bcl_performance(path, int(item["vf"]))
        if len(rows) != contract["bcl"]["expected_blocks_per_file"]:
            raise RuntimeError(f"unexpected BCL block count for {path}: {len(rows)}")
        performance_rows.extend(rows)
        cell_lineage.extend(cells)

    vor_item = contract["voronoi"]["primary_performance_file"]
    vor_path = checked(vor_item, "voronoi_primary")
    vor_rows, vor_cells = extract_voronoi_performance(vor_path, vor_item["sheet"])
    if len(vor_rows) != contract["voronoi"]["expected_rows"]:
        raise RuntimeError(f"unexpected Voronoi row count: {len(vor_rows)}")
    performance_rows.extend(vor_rows)
    cell_lineage.extend(vor_cells)
    checked(contract["voronoi"]["worked_crosscheck_file"], "voronoi_worked_crosscheck")

    relevant_x = [row for row in structural_rows if row["source_group"] in {"bcl", "voronoi"}]
    y_by_key = {row["model_key"]: row for row in performance_rows}
    x_by_key = {row["model_key"]: row for row in relevant_x}
    if len(y_by_key) != len(performance_rows):
        raise RuntimeError("duplicate performance model_key")
    if len(x_by_key) != len(relevant_x):
        raise RuntimeError("duplicate structural model_key")

    target_names = contract["target_policy"]["targets"]
    target_registry: list[dict] = []
    for target in target_names:
        values = [row[f"y__{target}"] for row in performance_rows]
        finite = sum(value is not None for value in values)
        target_registry.append(
            {
                "target_id": TARGET_IDS[target],
                "target_name_raw": target,
                "test_domain": "compression",
                "status": "available_unselected" if finite else "unavailable",
                "finite_count": finite,
                "missing_count": len(values) - finite,
                "objective_direction": "unresolved",
                "unit": "unresolved",
            }
        )

    all_keys = sorted(set(y_by_key) | set(x_by_key))
    row_registry: list[dict] = []
    joined_rows: list[dict] = []
    for key in all_keys:
        y_row = y_by_key.get(key)
        x_row = x_by_key.get(key)
        if y_row and x_row:
            status = "eligible_exact_model_label_join"
        elif y_row:
            status = "y_only_missing_structural_factors"
        else:
            status = "x_only_missing_performance"
        row_registry.append(
            {
                "model_key": key,
                "structural_model_id": x_row.get("structural_model_id", "") if x_row else "",
                "raw_performance_id": y_row.get("raw_model_id", "") if y_row else "",
                "source_group": (y_row or x_row)["source_group"],
                "family_group": (y_row or x_row)["family_group"],
                "vf_group": y_row.get("vf_group", "") if y_row else "",
                "crosswalk_status": status,
            }
        )
        if y_row and x_row:
            joined_rows.append({**x_row, **y_row, "crosswalk_status": status})

    x_headers = ["x__" + row["feature_name_raw"] for row in feature_registry]
    y_headers = ["y__" + target for target in target_names]
    performance_fields = [
        "model_key", "raw_model_id", "source_group", "family_group", "vf_group",
        "weight", "measured_vf", "height", "performance_source_file",
        "performance_sheet", "source_block_index", *y_headers,
    ]
    structural_fields = [
        "model_key", "structural_model_id", "source_group", "family_group",
        "structural_excel_row", *x_headers,
    ]
    joined_fields = [
        "model_key", "structural_model_id", "raw_model_id", "source_group",
        "family_group", "vf_group", "weight", "measured_vf", "height",
        "performance_source_file", "performance_sheet", "source_block_index",
        "structural_excel_row", "crosswalk_status", *x_headers, *y_headers,
    ]

    write_csv(output / "performance_y_wide_all.csv", performance_rows, performance_fields)
    write_csv(output / "structural_x_wide_relevant.csv", relevant_x, structural_fields)
    write_csv(output / "xy_eligible_exact_join.csv", joined_rows, joined_fields)
    write_csv(output / "row_registry.csv", row_registry)
    write_csv(output / "target_registry.csv", target_registry)
    write_csv(output / "feature_registry.csv", feature_registry)
    write_csv(output / "cell_lineage.csv", cell_lineage)
    write_csv(output / "source_audit.csv", source_audit)

    status_counts: dict[str, int] = {}
    for row in row_registry:
        status_counts[row["crosswalk_status"]] = status_counts.get(row["crosswalk_status"], 0) + 1
    joined_by_group: dict[str, int] = {}
    for row in joined_rows:
        joined_by_group[row["source_group"]] = joined_by_group.get(row["source_group"], 0) + 1

    qa = [
        {"check_id": "Q01", "check": "source hashes match", "passed": all(row["status"] == "PASS" for row in source_audit), "observed": f"{sum(row['status'] == 'PASS' for row in source_audit)}/{len(source_audit)}"},
        {"check_id": "Q02", "check": "BCL performance blocks 198", "passed": sum(row["source_group"] == "bcl" for row in performance_rows) == 198, "observed": sum(row["source_group"] == "bcl" for row in performance_rows)},
        {"check_id": "Q03", "check": "Voronoi performance rows 30", "passed": len(vor_rows) == 30, "observed": len(vor_rows)},
        {"check_id": "Q04", "check": "relevant structural rows 170", "passed": len(relevant_x) == 170, "observed": len(relevant_x)},
        {"check_id": "Q05", "check": "exact label joins BCL 137", "passed": joined_by_group.get("bcl", 0) == 137, "observed": joined_by_group.get("bcl", 0)},
        {"check_id": "Q06", "check": "exact ID joins Voronoi 30", "passed": joined_by_group.get("voronoi", 0) == 30, "observed": joined_by_group.get("voronoi", 0)},
        {"check_id": "Q07", "check": "16 target registry rows", "passed": len(target_registry) == 16, "observed": len(target_registry)},
        {"check_id": "Q08", "check": "all eligible joined target values finite", "passed": all(row[f"y__{target}"] is not None for row in joined_rows for target in target_names), "observed": sum(row[f"y__{target}"] is None for row in joined_rows for target in target_names)},
        {"check_id": "Q09", "check": "crosswalk is explicit, not row-order-only", "passed": all(row["crosswalk_status"] == "eligible_exact_model_label_join" for row in joined_rows), "observed": "model_key exact"},
        {"check_id": "Q10", "check": "no feature promoted", "passed": all(row["status"] != "promoted" for row in feature_registry), "observed": "unselected/excluded"},
    ]
    write_csv(output / "producer_QA.csv", qa)
    qa_pass = all(bool(row["passed"]) for row in qa)

    completed = datetime.now(KST)
    report = f"""# COMP-FACTORY-003 B/C/L + Voronoi intake

- Run ID: `{args.run_id}`
- Status: `{'PASS' if qa_pass else 'FAIL'}`
- Runtime: `{sys.executable}` / Python `{platform.python_version()}`
- B/C/L performance rows: `{sum(row['source_group'] == 'bcl' for row in performance_rows)}`
- Voronoi performance rows: `{len(vor_rows)}`
- Relevant structural rows: `{len(relevant_x)}`
- Exact x-y joins: `{len(joined_rows)}` (`B/C/L={joined_by_group.get('bcl', 0)}`, `Voronoi={joined_by_group.get('voronoi', 0)}`)
- Crosswalk statuses: `{json.dumps(status_counts, ensure_ascii=False, sort_keys=True)}`

## Source decision

`New_VF30/45/60.xlsx` is the primary B/C/L performance source because it contains explicit B1–T17 labels and populated cached formula outputs. `Compression_Worked_VF30/45/60.xlsx` and the small summary workbooks are retained as reference-only lineage assets; their generic trailing rows are not used as the current B/C/L model crosswalk.

`Voronoi_Summary.xlsx / Raw data` provides 30 explicit identifiers and is used as the Voronoi performance source. The worked workbook is retained for cross-check provenance.

## Claim boundary

This run performs source audit, explicit ID normalization, x-y joining, and QA only. It does not promote features, select a final model, claim production prediction, or claim inverse-design success.
"""
    (output / "REPORT.md").write_text(report, encoding="utf-8")

    manifest_files = []
    for path in sorted(output.iterdir()):
        if path.is_file():
            manifest_files.append({"file": path.name, "bytes": path.stat().st_size, "sha256": sha256(path)})
    manifest = {
        "run_id": args.run_id,
        "status": "PASS" if qa_pass else "FAIL",
        "started_at_kst": started.isoformat(),
        "completed_at_kst": completed.isoformat(),
        "runtime": {"sys_executable": sys.executable, "python": platform.python_version()},
        "contract": {"path": str(contract_path), "sha256": sha256(contract_path)},
        "counts": {
            "performance_rows": len(performance_rows),
            "structural_rows_relevant": len(relevant_x),
            "eligible_exact_joins": len(joined_rows),
            "target_count": len(target_registry),
            "feature_registry_count": len(feature_registry),
            "crosswalk_status": status_counts,
        },
        "outputs": manifest_files,
    }
    (output / "MANIFEST.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output), "status": manifest["status"], "counts": manifest["counts"]}, ensure_ascii=False, indent=2))
    if not qa_pass:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
