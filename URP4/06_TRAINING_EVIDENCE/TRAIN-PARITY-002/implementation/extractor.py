from __future__ import annotations

import json
import math
import re
from collections import Counter
from dataclasses import asdict
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from .contracts import ExtractionContract, ROUND_X_DECIMALS, SOURCE_SHA256, WORKBOOK_SHA256, target_policy, validate_target_policy
from .hashing import canonical_hash, verify_file_hash


def _ffill_horiz(values):
    out, last = [], None
    for value in values:
        if value is None or str(value).strip() == "":
            out.append(last)
        else:
            last = str(value).strip()
            out.append(last)
    return out


def _sanitize_header(text):
    text = "" if text is None else str(text)
    text = text.replace("\n", " ").replace("\r", " ").strip()
    return re.sub(r"\s+", " ", text)


def _make_unique(names):
    seen, out = Counter(), []
    for name in names:
        base = name if name else "Unnamed"
        seen[base] += 1
        out.append(base if seen[base] == 1 else f"{base}__{seen[base]}")
    return out


def build_headers(ws, start_letter, end_letter, main_row, sub_rows):
    start = column_index_from_string(start_letter)
    end = column_index_from_string(end_letter)
    columns = list(range(start, end + 1))
    row_values = []
    for row in [main_row, *sub_rows]:
        row_values.append(_ffill_horiz([ws.cell(row=row, column=column).value for column in columns]))
    headers = []
    for position, column in enumerate(columns):
        parts = []
        for values in row_values:
            value = values[position]
            if value is not None and str(value).strip() != "":
                value = _sanitize_header(value)
                if not parts or parts[-1] != value:
                    parts.append(value)
        headers.append(" | ".join(parts or [f"COL_{get_column_letter(column)}"]))
    return _make_unique(headers), [get_column_letter(column) for column in columns]


def _extract(ws, start_letter, end_letter, start_row, end_row, contract):
    headers, letters = build_headers(ws, start_letter, end_letter, contract.header_main_row, contract.header_sub_rows)
    start = column_index_from_string(start_letter)
    end = column_index_from_string(end_letter)
    data = [[ws.cell(row=row, column=column).value for column in range(start, end + 1)] for row in range(start_row, end_row + 1)]
    frame = pd.DataFrame(data, columns=headers)
    frame.attrs["excel_letters"] = letters
    return frame


def _safe_numeric(frame):
    result = frame.copy()
    for column in result.columns:
        result[column] = pd.to_numeric(result[column], errors="coerce")
    return result


def _json_scalar(value):
    if value is None or (isinstance(value, float) and math.isnan(value)) or pd.isna(value):
        return None
    if isinstance(value, (np.integer, int)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return float(value)
    return str(value)


def _series_hash(series):
    return canonical_hash([_json_scalar(value) for value in series.tolist()])


def _mask_hash(series):
    return canonical_hash([bool(value) for value in series.isna().tolist()])


def build_group_ids(x_frame: pd.DataFrame):
    rounded = x_frame.round(ROUND_X_DECIMALS)
    x_key = rounded.astype(str).agg("||".join, axis=1)
    codes, _ = pd.factorize(x_key, sort=False)
    return pd.Series(codes, name="GROUP_ID"), x_key.rename("X_KEY")


def validate_header_order(actual_letters, expected_letters, label):
    if list(actual_letters) != list(expected_letters):
        raise ValueError(f"{label}_COLUMN_ORDER_MISMATCH")


def extract_dataset(source_path: Path, workbook_path: Path, contract: ExtractionContract | None = None):
    contract = contract or ExtractionContract()
    verify_file_hash(source_path, SOURCE_SHA256, "SOURCE")
    verify_file_hash(workbook_path, WORKBOOK_SHA256, "WORKBOOK")
    workbook = load_workbook(workbook_path, data_only=True, read_only=False)
    if contract.sheet not in workbook.sheetnames:
        raise ValueError(f"SHEET_MISMATCH: {contract.sheet!r} not found")
    ws = workbook[contract.sheet]
    x_raw = _extract(ws, contract.x_start, contract.x_end, contract.data_row_start, contract.data_row_end, contract)
    y_all = _extract(ws, contract.y_start, contract.y_end, contract.data_row_start, contract.data_row_end, contract)
    y_mapping = dict(zip(y_all.attrs["excel_letters"], y_all.columns))
    missing_targets = [column for column in contract.target_columns if column not in y_mapping]
    if missing_targets:
        raise ValueError(f"TARGET_COLUMN_MISMATCH: {missing_targets}")
    y_raw = y_all[[y_mapping[column] for column in contract.target_columns]].copy()
    y_raw.attrs["excel_letters"] = list(contract.target_columns)
    if len(x_raw) != contract.expected_rows or len(y_raw) != contract.expected_rows:
        raise ValueError(f"ROW_COUNT_MISMATCH: X={len(x_raw)} Y={len(y_raw)}")
    x_numeric = _safe_numeric(x_raw).loc[:, lambda frame: frame.notna().any(axis=0)].copy()
    y_numeric = _safe_numeric(y_raw).loc[:, lambda frame: frame.notna().any(axis=0)].copy()
    if x_numeric.shape != (contract.expected_rows, contract.expected_x_features):
        raise ValueError(f"X_SHAPE_MISMATCH: {x_numeric.shape}")
    if y_numeric.shape != (contract.expected_rows, contract.expected_targets):
        raise ValueError(f"Y_SHAPE_MISMATCH: {y_numeric.shape}")
    expected_x_letters = [get_column_letter(index) for index in range(column_index_from_string(contract.x_start), column_index_from_string(contract.x_end) + 1)]
    validate_header_order(x_raw.attrs["excel_letters"], expected_x_letters, "X")
    validate_header_order(y_raw.attrs["excel_letters"], contract.target_columns, "Y_SELECTED")
    group_id, x_key = build_group_ids(x_numeric)
    policy = target_policy()
    validate_target_policy(policy)
    return {
        "contract": contract,
        "x_raw": x_raw,
        "y_raw": y_raw,
        "x": x_numeric,
        "y": y_numeric,
        "x_letters": expected_x_letters,
        "y_letters": list(contract.target_columns),
        "group_id": group_id,
        "x_key": x_key,
        "target_policy": policy,
        "source_rows": list(range(contract.data_row_start, contract.data_row_end + 1)),
    }


def write_dataset_artifacts(dataset, output_dir: Path):
    output_dir.mkdir(parents=True, exist_ok=True)
    x, y = dataset["x"], dataset["y"]
    row_rows = []
    for position, excel_row in enumerate(dataset["source_rows"]):
        x_row, y_row = x.iloc[position], y.iloc[position]
        row_rows.append({
            "row_id": f"EXCEL-ROW-{excel_row:04d}",
            "row_position": position,
            "original_excel_row": excel_row,
            "x_raw_value_hash": _series_hash(x_row),
            "y_raw_value_hash": _series_hash(y_row),
            "x_nan_mask_hash": _mask_hash(x_row),
            "y_nan_mask_hash": _mask_hash(y_row),
            "x_key": dataset["x_key"].iloc[position],
            "x_key_hash": canonical_hash(dataset["x_key"].iloc[position]),
            "group_id": int(dataset["group_id"].iloc[position]),
        })
    pd.DataFrame(row_rows).to_csv(output_dir / "ROW_IDENTITY_LEDGER.csv", index=False, encoding="utf-8-sig")
    column_rows = []
    for role, frame, letters in (("feature", x, dataset["x_letters"]), ("target", y, dataset["y_letters"])):
        for position, (column, letter) in enumerate(zip(frame.columns, letters)):
            column_rows.append({
                "role": role,
                "ordinal": position,
                "excel_col": letter,
                "normalized_header": column,
                "dtype": str(frame[column].dtype),
                "nan_count": int(frame[column].isna().sum()),
                "raw_value_hash": _series_hash(frame[column]),
                "nan_mask_hash": _mask_hash(frame[column]),
            })
    pd.DataFrame(column_rows).to_csv(output_dir / "COLUMN_IDENTITY_LEDGER.csv", index=False, encoding="utf-8-sig")
    target_rows = []
    for policy, column in zip(dataset["target_policy"], y.columns):
        mask = y[column].notna()
        target_rows.append({
            **policy,
            "normalized_header": column,
            "sample_count": int(mask.sum()),
            "group_count": int(dataset["group_id"][mask].nunique()),
            "target_value_hash": _series_hash(y.loc[mask, column]),
            "target_nan_mask_hash": _mask_hash(y[column]),
        })
    pd.DataFrame(target_rows).to_csv(output_dir / "TARGET_SAMPLE_LEDGER.csv", index=False, encoding="utf-8-sig")
    manifest = {
        "execution_status": "no_fit_no_prediction",
        "contract": asdict(dataset["contract"]),
        "rows": len(x),
        "x_features": x.shape[1],
        "targets": y.shape[1],
        "groups": int(dataset["group_id"].nunique()),
        "row_ledger": "ROW_IDENTITY_LEDGER.csv",
        "column_ledger": "COLUMN_IDENTITY_LEDGER.csv",
        "target_ledger": "TARGET_SAMPLE_LEDGER.csv",
    }
    (output_dir / "DATASET_MANIFEST.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest
