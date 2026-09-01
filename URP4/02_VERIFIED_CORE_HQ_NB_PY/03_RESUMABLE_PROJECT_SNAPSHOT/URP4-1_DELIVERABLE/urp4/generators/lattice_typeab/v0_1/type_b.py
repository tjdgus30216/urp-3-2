"""Identity-locked Type-B workbook adapter.

The professor notebook refers to an author-local ``Variables.xlsx``. CINT-04 can
test the schema with a synthetic fixture, but refuses official Type-B generation
unless the caller supplies both a real path and its expected SHA-256.
"""

from __future__ import annotations

import math
from pathlib import Path
from typing import Sequence

import numpy as np
from openpyxl import load_workbook

from urp4.contracts.v0_1.canonical import sha256_file

from .graph import _key3, assign_uniform_radius_for_vf, deduplicate_nodes_edges, ensure_connected, tile_unit_cell
from .models import GenerationConfig, LatticeGraph, TypeBInputUnavailable


TEMPLATE_COLUMNS = (
    "Start-x", "Start-y", "Start-z", "Start-Radius",
    "End-x", "End-y", "End-z", "End-radius",
)


def _number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _build_graph(rows: list[tuple[float, ...]], source_name: str, config: GenerationConfig) -> LatticeGraph | None:
    point_to_id: dict[tuple[float, float, float], int] = {}
    nodes: list[list[float]] = []
    edges: list[tuple[int, int]] = []
    radii: list[float] = []

    def node_id(point: Sequence[float]) -> int:
        key = _key3(np.asarray(point), config.merge_round_digits)
        if key not in point_to_id:
            point_to_id[key] = len(nodes)
            nodes.append([float(point[0]), float(point[1]), float(point[2])])
        return point_to_id[key]

    for sx, sy, sz, sr, ex, ey, ez, er in rows:
        first, second = node_id((sx, sy, sz)), node_id((ex, ey, ez))
        if first == second:
            continue
        edges.append(tuple(sorted((first, second))))
        radii.append(max(1e-6, 0.5 * (sr + er)))
    if len(nodes) < 2 or not edges:
        return None
    node_array, edge_array = deduplicate_nodes_edges(
        np.asarray(nodes), edges, config.merge_round_digits, config.min_edge_length_mm
    )
    radius = float(np.mean(radii)) if radii else 0.5
    return LatticeGraph(
        node_array, edge_array, np.full(len(edge_array), radius), "B", math.nan, math.nan, source_name
    )


def load_type_b_workbook(
    path: str | Path | None,
    expected_sha256: str | None,
    config: GenerationConfig,
) -> list[LatticeGraph]:
    if path is None or expected_sha256 is None:
        raise TypeBInputUnavailable("Type-B requires an explicit workbook path and expected SHA-256")
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise TypeBInputUnavailable(f"Type-B workbook does not exist: {workbook_path}")
    observed = sha256_file(workbook_path)
    if observed != expected_sha256.lower():
        raise TypeBInputUnavailable(f"Type-B workbook SHA-256 mismatch: expected {expected_sha256}, got {observed}")
    imported: list[LatticeGraph] = []
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        for sheet in workbook.worksheets:
            rows: list[tuple[float, ...]] = []
            header = tuple(sheet.cell(1, column).value for column in range(1, 9))
            if header == TEMPLATE_COLUMNS:
                for raw in sheet.iter_rows(min_row=2, values_only=True):
                    values = tuple(_number(item) for item in raw[:8])
                    if all(item is not None for item in values):
                        rows.append(tuple(float(item) for item in values))
            else:
                third = tuple(
                    str(sheet.cell(3, column).value).strip() if sheet.cell(3, column).value is not None else ""
                    for column in range(1, 14)
                )
                if third[1:4] == ("x", "y", "z") and third[4:7] == ("S-x", "S-y", "S-z") and third[7:10] == ("E-x", "E-y", "E-z"):
                    for row_index in range(4, sheet.max_row + 1):
                        xyz = tuple(_number(sheet.cell(row_index, column).value) for column in range(5, 11))
                        if all(item is not None for item in xyz):
                            sx, sy, sz, ex, ey, ez = (float(item) for item in xyz)
                            rows.append((sx, sy, sz, 0.5, ex, ey, ez, 0.5))
            graph = _build_graph(rows, f"{workbook_path.stem}:{sheet.title}", config)
            if graph is not None:
                imported.append(graph)
    finally:
        workbook.close()
    if not imported:
        raise TypeBInputUnavailable("Type-B workbook contains no supported non-empty lattice sheet")
    return imported


def normalize_or_tile_imported_lattice(graph: LatticeGraph, config: GenerationConfig) -> tuple[np.ndarray, np.ndarray]:
    nodes, edges = np.asarray(graph.nodes).copy(), np.asarray(graph.edges).copy()
    maximum_span = float(np.max(nodes.max(axis=0) - nodes.min(axis=0)))
    if maximum_span <= config.cell_size_mm * 1.35:
        return tile_unit_cell(nodes, edges, config)
    center = 0.5 * (nodes.max(axis=0) + nodes.min(axis=0))
    nodes -= center
    nodes *= config.total_length_mm / maximum_span if maximum_span > 1e-9 else 1.0
    half = config.total_length_mm / 2.0
    nodes = np.clip(nodes, -half, half)
    nodes, edges = deduplicate_nodes_edges(nodes, edges, config.merge_round_digits, config.min_edge_length_mm)
    return nodes, ensure_connected(nodes, edges)


def create_type_b_model(
    target_vf: float, base_graph: LatticeGraph, config: GenerationConfig, rng: np.random.Generator
) -> LatticeGraph:
    nodes, edges = normalize_or_tile_imported_lattice(base_graph, config)
    half = config.total_length_mm / 2.0
    boundary = np.any(np.isclose(np.abs(nodes), half, atol=1e-6), axis=1)
    if config.type_b_node_jitter_mm > 0:
        jitter = rng.normal(0.0, config.type_b_node_jitter_mm, size=nodes.shape)
        jitter[boundary] *= 0.25
        nodes += jitter
    if config.type_b_anisotropic_scale_std > 0:
        nodes *= rng.normal(1.0, config.type_b_anisotropic_scale_std, size=3)
    nodes = np.clip(nodes, -half, half)
    if config.type_b_edge_dropout_prob > 0 and len(edges) > 4:
        keep = rng.random(len(edges)) >= config.type_b_edge_dropout_prob
        if keep.sum() >= 4:
            edges = edges[keep]
    nodes, edges = deduplicate_nodes_edges(nodes, edges, config.merge_round_digits, config.min_edge_length_mm)
    edges = ensure_connected(nodes, edges)
    radii, actual_vf = assign_uniform_radius_for_vf(nodes, edges, target_vf, config.total_length_mm)
    result = LatticeGraph(nodes, edges, radii, "B", target_vf, actual_vf, base_graph.source_name)
    result.validate()
    return result

